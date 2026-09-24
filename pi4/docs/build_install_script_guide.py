"""install.shのセットアップ手順ドキュメント(README.md相当)をDOCXへ変換する。
README.mdの簡潔な構成をそのまま踏襲する(詳細な逐次解説は
install_script_guide.mdを参照)。build_operation_manual.pyと同じ体裁
(見出し色・表紙)を踏襲する。
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "install_script_guide.docx"

d = Document()
sec = d.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)

styles = d.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25
for name, size, color, before, after in (
    ("Heading 1", 16, "2E74B5", 18, 10),
    ("Heading 2", 13, "2E74B5", 14, 7),
    ("Heading 3", 12, "1F4D78", 10, 5),
):
    s = styles[name]
    s.font.name = "Calibri"
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)

# 表紙
p = d.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(140)
r = p.add_run("Shonan_Lite for RasPI4")
r.bold = True
r.font.size = Pt(26)
r.font.color.rgb = RGBColor(31, 77, 120)
p = d.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("セットアップガイド")
r.bold = True
r.font.size = Pt(18)
p = d.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(24)
p.add_run(
    "対象: pi4/scripts/install.sh\n"
    "新規Pi4への一括セットアップ手順(README.md相当)\n"
    "版: 2026-08-15"
).font.size = Pt(11)
d.add_page_break()

d.add_heading("目次", level=1)


def add_toc_entry(text: str, level: int = 1) -> None:
    para = d.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.3 * (level - 1))
    r = para.add_run(text)
    if level == 1:
        r.bold = True


TOC_ENTRIES = [
    ("概要", 1),
    ("0. Raspberry Pi OSのインストール", 1),
    ("1. install.shの実行", 1),
    ("実行条件(前提条件)", 1),
    ("関連ドキュメント", 1),
]
for _text, _level in TOC_ENTRIES:
    add_toc_entry(_text, _level)
d.add_page_break()


def add_code_block(lines: list[str]) -> None:
    table = d.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    cell.paragraphs[0].text = ""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    cell._tc.get_or_add_tcPr().append(shading)
    first = True
    for line in lines:
        para = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        run = para.add_run(line if line else " ")
        run.font.name = "Courier New"
        run.font.size = Pt(9.5)
        para.paragraph_format.space_after = Pt(0)
    d.add_paragraph()


def add_table(headers: list[str], rows: list[list[str]]) -> None:
    table = d.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for para in hdr_cells[i].paragraphs:
            for run in para.runs:
                run.font.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    d.add_paragraph()


def add_para(text: str, bold_prefix: str | None = None) -> None:
    para = d.add_paragraph()
    if bold_prefix:
        r = para.add_run(bold_prefix)
        r.bold = True
        para.add_run(text)
    else:
        para.add_run(text)


def add_label(text: str) -> None:
    para = d.add_paragraph()
    r = para.add_run(text)
    r.bold = True


def add_bullets(items: list[str]) -> None:
    for item in items:
        d.add_paragraph(item, style="List Bullet")


def add_numbered(items: list[str]) -> None:
    for item in items:
        d.add_paragraph(item, style="List Number")


# 本文 ---------------------------------------------------------------------

d.add_heading("概要", level=1)
add_para(
    "Raspberry Pi 4 + ADALM-Pluto+によるDVB-S2 DATV送受信タッチGUIシステム。"
    "本体はpi4/配下(Python/PyQt5、eglfs直描画)。同系統の別プラットフォーム"
    "移植版(Android/iOS)は別リポジトリ(Shonan_Lite-android/Shonan_Lite-iPad等)"
    "で管理している。"
)

d.add_heading("0. Raspberry Pi OSのインストール", level=1)
add_para("Pi4本体に、あらかじめRaspberry Pi OSをインストールしておく。")
add_numbered([
    "PCでRaspberry Pi Imager(https://www.raspberrypi.com/software/)を起動する。",
    "「デバイスを選択」でRaspberry Pi 4を選ぶ。",
    "「OSを選択」でRaspberry Pi OS (64-bit)を選ぶ(32bit版は不可。"
    "「実行条件」参照)。",
    "「ストレージを選択」で書き込み先のmicroSD/NVMeを選ぶ。",
    "歯車アイコン(詳細設定)で、ホスト名・ユーザー名/パスワード・Wi-Fi・"
    "SSH有効化を事前設定しておくと、初回起動後すぐSSH接続できる。",
    "「書き込む」を実行し、完了後microSD/NVMeをPi4に取り付けて起動する。",
    "PCからssh <ユーザー名>@<ホスト名>.localで接続できることを確認する。",
])

d.add_heading("1. install.shの実行", level=1)
add_para("SSH接続したPi4上で:")
add_code_block([
    "git clone https://github.com/kazushinjo/Shonan_lite-PI4.git",
    "cd Shonan_lite-PI4",
    "./pi4/scripts/install.sh          # HTTPSでclone(既定)",
    "# ./pi4/scripts/install_ssh.sh    # GitHubにSSH鍵を登録済みならこちらでも可",
])
add_para(
    "日本語入力ビルド・受信(RX)用GNU Radio/gr-dvbs2rxビルドはそれぞれ省略して"
    "時間短縮できる(RX用を省略すると受信機能は使えなくなる):"
)
add_code_block([
    "SKIP_JA_KEYBOARD=1 SKIP_GNURADIO_BUILD=1 ./pi4/scripts/install.sh",
])
add_para(
    "完了後、shonan-gui.serviceがsystemdに登録されGUIが自動起動する。あわせて"
    "/boot/firmware/config.txtへavoid_warnings=1(電源電圧警告アイコンの表示抑制)"
    "を未設定なら自動で追記する(反映には再起動が必要)。"
)

d.add_heading("実行条件(前提条件)", level=1)
add_para("一般ユーザーが実行してinstall.shが正常完了するには、以下が必要。")
add_bullets([
    "Raspberry Pi OS 64bit(aarch64)であること(32bit版不可)",
    "gitが事前にインストール済みであること",
    "GitHub/apt配布ミラーへのインターネット到達性",
    "sudoが使える対話的な実行(パスワード入力に応答できるtty)",
    "patchコマンドが使えること",
])
add_para(
    "詳細な各手順の解説・トラブルシュートはpi4/docs/install_script_guide.mdを"
    "参照。"
)

d.add_heading("関連ドキュメント", level=1)
add_bullets([
    "pi4/docs/install_script_guide.md — install.shの詳細ガイド",
    "pi4/docs/qtvirtualkeyboard_ja_build.md — 日本語オンスクリーンキーボードの"
    "ビルド手順・ハマりどころ",
    "pi4/docs/shonan_pi4_operation_manual.docx / pi4/gui/manual_content.py — "
    "GUIの操作説明書(アプリ内Helpと同内容)",
    "pi4/third_party/rpi-dvbs2-receiver-gui/ — GNU Radio/gr-dvbs2rx受信フロー"
    "グラフの参考実装(kazushinjo/rpi-dvbs2-receiver-guiより取り込み)",
])

d.save(OUT)
print(f"saved: {OUT}")
