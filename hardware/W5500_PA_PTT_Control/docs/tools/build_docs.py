"""hardware/W5500_PA_PTT_Control/docsのMarkdown文書から、Word版(.docx)とPDF版(.pdf)を作る。

使い方(macOSで実行。pandoc・Word・LibreOfficeは不要。python-docxが必要):
    python3 hardware/W5500_PA_PTT_Control/docs/tools/build_docs.py

Markdownを見出し・段落・表・箇条書き・コード・引用の部品に分け、Word版はpython-docxで、
PDF版はHTMLにしてから同じフォルダのhtml2pdf.swift(macOS標準の文書描画機能で印刷する)で
作る。対応している書式は、このフォルダの文書で使っているもの(見出し・表・箇条書き・
番号付きリスト・コード・引用・太字・インラインコード・リンク・front matter)だけ。
"""
from __future__ import annotations

import html
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DOCS_DIR = Path(__file__).resolve().parents[1]
TOOLS_DIR = Path(__file__).resolve().parent
# Word版・PDF版を作る文書(日本語版のみ。英語版はMarkdownのみで管理する)。
TARGETS = ["W5500_PA_PTT_Control_仕様書.md", "MCU1_J1_W5500_接続一覧.md"]

CSS = """
body { font-family: 'Hiragino Sans', 'Helvetica Neue', sans-serif; font-size: 10.5pt; line-height: 1.5; }
h1 { font-size: 17pt; color: #1F4D78; } h2 { font-size: 13.5pt; color: #2E74B5; }
h3 { font-size: 11.5pt; color: #2E74B5; }
table { border-collapse: collapse; margin: 6px 0; }
th, td { border: 1px solid #888; padding: 3px 6px; vertical-align: top; font-size: 9.5pt; }
th { background: #E8EEF6; }
pre { font-family: Menlo, monospace; font-size: 8.5pt; background: #F4F4F4; border: 1px solid #ccc; padding: 6px; }
code { font-family: Menlo, monospace; font-size: 9pt; }
blockquote { border-left: 3px solid #bbb; margin-left: 0; padding-left: 10px; color: #444; }
"""


# ---------------------------------------------------------------- Markdownの解析

def join_lines(lines: list[str]) -> str:
    """改行で分かれた文をつなぐ。前後がどちらも英数字記号のときだけ空白を入れる
    (日本語の文の途中に余計な空白が入らないように)。"""
    out = ""
    for line in (l.strip() for l in lines):
        if out and line and ord(out[-1]) < 0x2E80 and ord(line[0]) < 0x2E80:
            out += " "
        out += line
    return out


def split_row(line: str) -> list[str]:
    cells = re.split(r"(?<!\\)\|", line.strip().strip("|"))
    return [c.strip().replace("\\|", "|") for c in cells]


LIST_RE = re.compile(r"^(\s*)([-*]|\d+\.)\s+(.*)")


def parse(md: str) -> list[tuple]:
    """Markdownを部品のリストにする。

    ("h", level, text) / ("p", text) / ("table", header, rows) / ("list", ordered, [(level, text)])
    / ("code", text) / ("quote", [text, ...]) / ("hr",)
    """
    lines = md.splitlines()
    if lines and lines[0].strip() == "---":  # front matter
        lines = lines[lines.index("---", 1) + 1:]
    blocks: list[tuple] = []
    para: list[str] = []

    def flush():
        if para:
            blocks.append(("p", join_lines(para)))
            para.clear()

    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            flush()
            j = i + 1
            while j < len(lines) and not lines[j].startswith("```"):
                j += 1
            blocks.append(("code", "\n".join(lines[i + 1:j])))
            i = j + 1
        elif m := re.match(r"^(#{1,4})\s+(.*)", line):
            flush()
            blocks.append(("h", len(m.group(1)), m.group(2)))
            i += 1
        elif line.strip() == "---":
            flush()
            blocks.append(("hr",))
            i += 1
        elif line.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            blocks.append(("table", split_row(rows[0]), [split_row(r) for r in rows[2:]]))
        elif line.startswith(">"):
            flush()
            quote = []
            while i < len(lines) and lines[i].startswith(">"):
                quote.append(lines[i][1:].strip())
                i += 1
            paras = [join_lines(p.split("\n")) for p in "\n".join(quote).split("\n\n") if p.strip()]
            blocks.append(("quote", paras))
        elif m := LIST_RE.match(line):
            flush()
            ordered = m.group(2)[0].isdigit()
            items: list[tuple[int, str]] = []
            while i < len(lines):
                mm = LIST_RE.match(lines[i])
                if mm:
                    items.append((1 if len(mm.group(1)) >= 2 else 0, mm.group(3)))
                elif lines[i].startswith("  ") and lines[i].strip() and items:
                    items[-1] = (items[-1][0], join_lines([items[-1][1], lines[i]]))
                else:
                    break
                i += 1
            blocks.append(("list", ordered, items))
        else:
            if line.strip():
                para.append(line)
            else:
                flush()
            i += 1
    flush()
    return blocks


INLINE_RE = re.compile(r"`([^`]*)`|\*\*(.+?)\*\*|\[([^\]]+)\]\(([^)]+)\)")


def spans(text: str, bold: bool = False) -> list[tuple[str, bool, bool]]:
    """インライン書式を(文字列, 太字, コード)の並びにする。太字やリンクの中のコードにも対応。"""
    out: list[tuple[str, bool, bool]] = []
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], bold, False))
        if m.group(1) is not None:
            out.append((m.group(1), bold, True))
        elif m.group(2) is not None:
            out.extend(spans(m.group(2), True))
        else:
            out.extend(spans(m.group(3), bold))  # リンクは文字だけ残す
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], bold, False))
    return out


# ---------------------------------------------------------------- HTML(PDF用)

def inline_html(text: str) -> str:
    parts = []
    for s, b, c in spans(text):
        s = html.escape(s)
        if c:
            s = f"<code>{s}</code>"
        if b:
            s = f"<b>{s}</b>"
        parts.append(s)
    return "".join(parts)


def to_html(blocks: list[tuple]) -> str:
    body = []
    for blk in blocks:
        kind = blk[0]
        if kind == "h":
            body.append(f"<h{blk[1]}>{inline_html(blk[2])}</h{blk[1]}>")
        elif kind == "p":
            body.append(f"<p>{inline_html(blk[1])}</p>")
        elif kind == "hr":
            body.append("<hr>")
        elif kind == "code":
            body.append(f"<pre>{html.escape(blk[1])}</pre>")
        elif kind == "quote":
            body.append("<blockquote>" + "".join(f"<p>{inline_html(p)}</p>" for p in blk[1]) + "</blockquote>")
        elif kind == "table":
            rows = ["<tr>" + "".join(f"<th>{inline_html(c)}</th>" for c in blk[1]) + "</tr>"]
            rows += ["<tr>" + "".join(f"<td>{inline_html(c)}</td>" for c in r) + "</tr>" for r in blk[2]]
            body.append("<table>" + "".join(rows) + "</table>")
        elif kind == "list":
            tag = "ol" if blk[1] else "ul"
            out, depth = [f"<{tag}>"], 0
            for level, text in blk[2]:
                while depth < level:
                    out.append("<ul>")
                    depth += 1
                while depth > level:
                    out.append("</ul>")
                    depth -= 1
                out.append(f"<li>{inline_html(text)}</li>")
            out += ["</ul>"] * depth + [f"</{tag}>"]
            body.append("".join(out))
    return ("<html><head><meta charset='utf-8'><style>" + CSS + "</style></head><body>"
            + "\n".join(body) + "</body></html>")


# ---------------------------------------------------------------- Word

def add_runs(paragraph, text: str, size: float | None = None) -> None:
    for s, b, c in spans(text):
        run = paragraph.add_run(s)
        run.bold = b or None
        if c:
            run.font.name = "Menlo"
            run.font.size = Pt(9)
        elif size:
            run.font.size = Pt(size)


def set_cell_shading(cell, color: str) -> None:
    shd = cell._tc.get_or_add_tcPr().makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:fill"): color})
    cell._tc.get_or_add_tcPr().append(shd)


def to_docx(blocks: list[tuple], out: Path) -> None:
    d = Document()
    normal = d.styles["Normal"]
    normal.font.name = "Hiragino Sans"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans")
    normal.font.size = Pt(10.5)
    for blk in blocks:
        kind = blk[0]
        if kind == "h":
            add_runs(d.add_heading("", level=min(blk[1], 3)), blk[2])
        elif kind == "p":
            add_runs(d.add_paragraph(), blk[1])
        elif kind == "hr":
            continue
        elif kind == "code":
            p = d.add_paragraph()
            run = p.add_run(blk[1])
            run.font.name = "Menlo"
            run.font.size = Pt(8.5)
        elif kind == "quote":
            for text in blk[1]:
                p = d.add_paragraph()
                p.paragraph_format.left_indent = Pt(14)
                add_runs(p, text)
                for run in p.runs:
                    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        elif kind == "table":
            header, rows = blk[1], blk[2]
            table = d.add_table(rows=1 + len(rows), cols=len(header))
            table.style = "Table Grid"
            for j, text in enumerate(header):
                cell = table.cell(0, j)
                add_runs(cell.paragraphs[0], text, 9.5)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                set_cell_shading(cell, "E8EEF6")
            for i, row in enumerate(rows, start=1):
                for j, text in enumerate(row[:len(header)]):
                    add_runs(table.cell(i, j).paragraphs[0], text, 9.5)
        elif kind == "list":
            ordered = blk[1]
            for level, text in blk[2]:
                style = "List Number" if ordered and level == 0 else ("List Bullet 2" if level else "List Bullet")
                add_runs(d.add_paragraph(style=style), text)
    d.save(out)


# ----------------------------------------------------------------

def main() -> None:
    swift = TOOLS_DIR / "html2pdf.swift"
    for name in TARGETS:
        md_path = DOCS_DIR / name
        if not md_path.exists():
            print(f"skip: {name} がありません", file=sys.stderr)
            continue
        blocks = parse(md_path.read_text(encoding="utf-8"))
        docx_path = md_path.with_suffix(".docx")
        pdf_path = md_path.with_suffix(".pdf")
        to_docx(blocks, docx_path)
        with tempfile.TemporaryDirectory() as tmp:
            html_path = Path(tmp) / "doc.html"
            html_path.write_text(to_html(blocks), encoding="utf-8")
            subprocess.run(["swift", str(swift), str(html_path), str(pdf_path)], check=True,
                           stderr=subprocess.DEVNULL)
        print(f"{docx_path.name} / {pdf_path.name} を作成しました")


if __name__ == "__main__":
    main()
